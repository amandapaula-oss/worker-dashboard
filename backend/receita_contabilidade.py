# -*- coding: utf-8 -*-
"""Relatorio "Receita Contabilidade": toda a receita por mes, com centro de lucro
(profit center), BU, PEP completo, ID do projeto, cliente e valor.

Gera um xlsx com 3 abas:
  1. Receita Detalhada  — uma linha por lancamento de receita
  2. Por Mes e BU       — matriz BU x mes
  3. Por Centro de Lucro— matriz centro de lucro x mes

O centro de lucro sai de `no_hierarquia` (ex.: "DC002 Dedicated Teams") e, quando a
linha nao tem, do cadastro mestre de PEPs do SAP (pep_master_sap.xlsx) pelo projeto.
"""
import io
import os
import re
from functools import lru_cache
from typing import List, Optional

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DIR = os.path.dirname(os.path.abspath(__file__))
MESTRE = os.path.join(DIR, "pep_master_sap.xlsx")

FONTE_BASE = "Aptos Narrow"
HDR_FILL = PatternFill("solid", fgColor="C04F15")
HDR_FONT = Font(name=FONTE_BASE, size=11, bold=True, color="FFFFFF")
CELL_FONT = Font(name=FONTE_BASE, size=11)
TOT_FILL = PatternFill("solid", fgColor="FFFFCC")
TOT_FONT = Font(name=FONTE_BASE, size=11, bold=True)
BORDA = Border(bottom=Side(style="thin", color="DDDDDD"))
NUM = "#.##0"

MES_PT = {1: "jan", 2: "fev", 3: "mar", 4: "abr", 5: "mai", 6: "jun",
          7: "jul", 8: "ago", 9: "set", 10: "out", 11: "nov", 12: "dez"}


def _rotulo_mes(periodo: str) -> str:
    try:
        ano, mes = str(periodo).split("-")[:2]
        return f"{MES_PT[int(mes)]}/{ano[2:]}"
    except Exception:
        return str(periodo)


@lru_cache(maxsize=1)
def _profit_center_por_pep() -> dict:
    """PEP (e raiz) -> codigo do centro de lucro, do cadastro mestre do SAP.
    Em cache: sao ~4 mil linhas de xlsx, nao vale reler a cada download.
    Falha silenciosa: o relatorio funciona sem o cadastro, so perde o preenchimento
    de quem nao tem no_hierarquia na propria linha."""
    mapa = {}
    if not os.path.exists(MESTRE):
        return mapa
    try:
        m = pd.read_excel(MESTRE, sheet_name="PEPs")
        m.columns = [str(c).strip() for c in m.columns]
        for _, r in m.iterrows():
            pc = str(r.get("profit_center") or "").strip().upper()
            if not re.match(r"^DC\d+$", pc):
                continue
            pep = str(r.get("pep") or "").strip().upper()
            if not pep:
                continue
            mapa.setdefault(pep, pc)
            mapa.setdefault(pep.split(".")[0], pc)
    except Exception as e:  # pragma: no cover
        print(f"[receita_contabilidade] cadastro mestre nao lido: {e}")
    return mapa


def _estiliza(ws, n_cols, larguras, linha_total=None, freeze="A5"):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.freeze_panes = freeze
    for j, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    if linha_total:
        for j in range(1, n_cols + 1):
            ws.cell(linha_total, j).fill = TOT_FILL
            ws.cell(linha_total, j).font = TOT_FONT


def gerar_xlsx_bytes(df: pd.DataFrame, periodos: Optional[List[str]] = None) -> bytes:
    """df: nova_base ja processada (e ja filtrada pelas BUs do usuario).
    periodos: lista de 'YYYY-MM' a incluir; None/vazio = todos."""
    d = df.copy()
    d["receita"] = pd.to_numeric(d.get("receita"), errors="coerce").fillna(0)
    d["periodo"] = d.get("periodo", "").astype(str).str.strip()
    fonte = d.get("fonte", pd.Series("", index=d.index)).fillna("").astype(str).str.strip()
    # Contabilidade = receita realizada: Budget e plano, "de para" e mapeamento auxiliar
    d = d[(d["receita"] != 0) & (~fonte.isin(["Budget", "de para"]))]
    if periodos:
        d = d[d["periodo"].isin([str(p).strip() for p in periodos])]
    if d.empty:
        # melhor um erro claro do que uma planilha vazia que parece dado
        raise ValueError("Não há receita para o período pedido.")

    pc_map = _profit_center_por_pep()
    pep = d.get("pep", pd.Series("", index=d.index)).fillna("").astype(str).str.strip()
    pep_base = d.get("pep_base", pd.Series("", index=d.index)).fillna("").astype(str).str.strip()
    nh = d.get("no_hierarquia", pd.Series("", index=d.index)).fillna("").astype(str).str.strip()

    cod = nh.str.extract(r"^(DC\d+)", expand=False)
    # quem nao tem DCxxx na linha: busca pelo PEP no cadastro mestre
    falta = cod.isna()
    if falta.any() and pc_map:
        alt = pep.str.upper().map(pc_map).fillna(pep_base.str.upper().map(pc_map))
        cod = cod.fillna(alt)
    nome_pc = nh.str.replace(r"^DC\d+\s*", "", regex=True).str.strip()

    det = pd.DataFrame({
        "Mês": d["periodo"].map(_rotulo_mes),
        "Competência": d["periodo"],
        "Empresa": d.get("empresa", "").fillna("").astype(str).str.strip(),
        "Centro de Lucro": cod.fillna(""),
        "Centro de Lucro (nome)": nome_pc,
        "BU": d.get("vertical", "").fillna("").astype(str).str.strip(),
        "PEP": pep,
        "ID Projeto": pep_base,
        "Cliente": d.get("nome_cliente", "").fillna("").astype(str).str.strip(),
        "Apuração": d.get("apuracao", "").fillna("").astype(str).str.strip(),
        "Fonte": fonte.loc[d.index],
        "Valor": d["receita"].round(2),
    }).sort_values(["Competência", "BU", "Cliente", "PEP"]).reset_index(drop=True)

    meses = sorted(det["Competência"].unique())
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        det.drop(columns=["Competência"]).to_excel(xw, sheet_name="Receita Detalhada",
                                                   index=False, startrow=3)
        ws = xw.sheets["Receita Detalhada"]
        ws["A1"] = "Receita Contabilidade — detalhe por lançamento"
        ws["A1"].font = Font(name=FONTE_BASE, size=12, bold=True)
        _tot = f"{det['Valor'].sum():,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        # deixa explicito o recorte: quem tem acesso a 1 BU precisa saber que o
        # arquivo nao e a empresa inteira
        _bus = sorted(b for b in det["BU"].unique() if str(b).strip())
        _rec = f" · BUs: {', '.join(_bus)}" if 0 < len(_bus) <= 3 else ""
        ws["A2"] = (f"{len(det)} lançamentos · {', '.join(_rotulo_mes(m) for m in meses)} · "
                    f"total R$ {_tot}{_rec} · receita realizada (exclui Budget)")
        ws["A2"].font = Font(name=FONTE_BASE, size=9, italic=True, color="666666")
        ncol = det.shape[1] - 1
        for j in range(1, ncol + 1):
            c = ws.cell(4, j)
            c.fill, c.font = HDR_FILL, HDR_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i in range(len(det)):
            for j in range(1, ncol + 1):
                c = ws.cell(5 + i, j)
                c.font, c.border = CELL_FONT, BORDA
                if j == ncol:
                    c.number_format = NUM
        tot = 5 + len(det)
        ws.cell(tot, 1, "TOTAL")
        ws.cell(tot, ncol, round(float(det["Valor"].sum()), 2)).number_format = NUM
        _estiliza(ws, ncol, [9, 15, 15, 24, 14, 22, 16, 34, 13, 20, 15], tot)

        # ---- matriz BU x mes ----
        piv = (det.pivot_table(index="BU", columns="Competência", values="Valor",
                               aggfunc="sum", fill_value=0)
               .reindex(columns=meses, fill_value=0))
        piv["Total"] = piv.sum(axis=1)
        piv = piv.sort_values("Total", ascending=False)
        piv.columns = [_rotulo_mes(c) if c != "Total" else "Total" for c in piv.columns]
        piv.reset_index().to_excel(xw, sheet_name="Por Mês e BU", index=False, startrow=3)
        ws2 = xw.sheets["Por Mês e BU"]
        ws2["A1"] = "Receita por BU e mês"
        ws2["A1"].font = Font(name=FONTE_BASE, size=12, bold=True)
        n2 = piv.shape[1] + 1
        for j in range(1, n2 + 1):
            c = ws2.cell(4, j)
            c.fill, c.font = HDR_FILL, HDR_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i in range(len(piv)):
            for j in range(1, n2 + 1):
                c = ws2.cell(5 + i, j)
                c.font, c.border = CELL_FONT, BORDA
                if j > 1:
                    c.number_format = NUM
        t2 = 5 + len(piv)
        ws2.cell(t2, 1, "TOTAL")
        for j in range(2, n2 + 1):
            col = piv.columns[j - 2]
            ws2.cell(t2, j, round(float(piv[col].sum()), 2)).number_format = NUM
        _estiliza(ws2, n2, [22] + [15] * (n2 - 1), t2)

        # ---- matriz centro de lucro x mes ----
        det["_pc"] = (det["Centro de Lucro"].replace("", "(sem centro de lucro)") + " " +
                      det["Centro de Lucro (nome)"]).str.strip()
        piv3 = (det.pivot_table(index="_pc", columns="Competência", values="Valor",
                                aggfunc="sum", fill_value=0)
                .reindex(columns=meses, fill_value=0))
        piv3["Total"] = piv3.sum(axis=1)
        piv3 = piv3.sort_values("Total", ascending=False)
        piv3.columns = [_rotulo_mes(c) if c != "Total" else "Total" for c in piv3.columns]
        piv3.index.name = "Centro de Lucro"
        piv3.reset_index().to_excel(xw, sheet_name="Por Centro de Lucro", index=False, startrow=3)
        ws3 = xw.sheets["Por Centro de Lucro"]
        ws3["A1"] = "Receita por centro de lucro e mês"
        ws3["A1"].font = Font(name=FONTE_BASE, size=12, bold=True)
        n3 = piv3.shape[1] + 1
        for j in range(1, n3 + 1):
            c = ws3.cell(4, j)
            c.fill, c.font = HDR_FILL, HDR_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i in range(len(piv3)):
            for j in range(1, n3 + 1):
                c = ws3.cell(5 + i, j)
                c.font, c.border = CELL_FONT, BORDA
                if j > 1:
                    c.number_format = NUM
        t3 = 5 + len(piv3)
        ws3.cell(t3, 1, "TOTAL")
        for j in range(2, n3 + 1):
            col = piv3.columns[j - 2]
            ws3.cell(t3, j, round(float(piv3[col].sum()), 2)).number_format = NUM
        _estiliza(ws3, n3, [30] + [15] * (n3 - 1), t3)

    buf.seek(0)
    return buf.getvalue()
