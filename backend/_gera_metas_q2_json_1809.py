# -*- coding: utf-8 -*-
"""Regera backend/metas_q2.json (18/09):
  Q2 -> 'RESUMO - 18.09.xlsx', aba 'Q2Y26 VIVO'  (numeros novos, pos-contabilidade)
  Q1 -> 'RESUMO.xlsx' original, aba 'Q1Y26 (Valor)' (congelado: o Q1 NAO e' recalculado,
        decisao da Amanda em 18/09 — o que foi pago fica como foi pago)
"""
import datetime, json, os, shutil, tempfile
import openpyxl

P = (r'C:\Users\amanda.paula\FCamara Consultoria e Formação\FCamara Files - CONTROLADORIA'
     r'\30. FP&A NOVO\15. Metas (Apuração)\2026\Metas Oficiais')
NOVO, VELHO = 'RESUMO - 18.09.xlsx', 'RESUMO.xlsx'


def abre(nome):
    """copia antes de abrir: o arquivo pode estar aberto no Excel da Amanda"""
    tmp = os.path.join(tempfile.gettempdir(), '_mq_' + nome)
    shutil.copy2(os.path.join(P, nome), tmp)
    return openpyxl.load_workbook(tmp, read_only=True, data_only=True)


def parse_tab(ws, periodo):
    out = []
    for r in range(4, ws.max_row + 1):
        bu, avaliado = ws.cell(r, 6).value, ws.cell(r, 7).value
        if not avaliado or str(avaliado).strip() in ('0', ''):
            continue
        def num(c):
            v = ws.cell(r, c).value
            try:
                return round(float(v), 6) if v is not None else None
            except (TypeError, ValueError):
                return None
        out.append({
            'periodo': periodo, 'contrato': ws.cell(r, 4).value,
            'bu': str(bu).strip() if bu else None, 'avaliado': str(avaliado).strip(),
            'posicao': ws.cell(r, 8).value, 'peso_meta': num(10),
            'atingimento_final': num(14), 'salario': num(15), 'quantidade': num(16),
            'apuracao_rs': num(17),
            'receita': {'realizado': num(19), 'meta': num(20), 'ating': num(21)},
            'mb': {'realizado': num(22), 'meta': num(23), 'delta_pp': num(24)},
            'lb': {'realizado': num(25), 'meta': num(26), 'ating': num(27)},
        })
    return out


wn, wv = abre(NOVO), abre(VELHO)
data = {
    'gerado_em': datetime.date.today().isoformat(),
    'fonte': ('Q2 = "RESUMO - 18.09.xlsx" aba Q2Y26 VIVO (base alinhada a contabilidade); '
              'Q1 = "RESUMO.xlsx" original aba Q1Y26 (Valor), congelado — o Q1 nao foi '
              'recalculado por decisao de 18/09'),
    'q2': parse_tab(wn['Q2Y26 VIVO'], 'Q2Y26'),
    'q1': parse_tab(wv['Q1Y26 (Valor)'], 'Q1Y26'),
}
wn.close(); wv.close()
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'metas_q2.json')
shutil.copy2(out, out.replace('.json', '_antes_1809.json'))
json.dump(data, open(out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(f'salvo {out}   (backup: metas_q2_antes_1809.json)')
print(f'Q2: {len(data["q2"])} avaliados | Q1: {len(data["q1"])}')
print(f'\n{"avaliado":30} {"BU":18} {"receita Q2":>15} {"bonus":>11}')
for x in sorted(data['q2'], key=lambda z: -(z['receita']['realizado'] or 0)):
    print(f"  {x['avaliado'][:28]:28} {str(x['bu'])[:18]:18} "
          f"{(x['receita']['realizado'] or 0):>15,.2f} {(x['apuracao_rs'] or 0):>11,.2f}")
