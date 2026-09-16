# -*- coding: utf-8 -*-
"""Confere o arquivo de Receita Contabilidade contra o P&L e olha a coerencia entre
os meses (saltos, linha que aparece/some, cliente que oscila)."""
import io, os
import openpyxl
import pandas as pd

os.environ.setdefault('SECRET_KEY', 'dbg-readonly')
os.environ.setdefault('ALLOWED_ORIGINS', '*')
from _supabase_creds import load_creds
load_creds()
import main
from receita_contabilidade import gerar_xlsx_bytes

MESES = ['2026-01', '2026-02', '2026-03', '2026-04', '2026-05', '2026-06']
# linha 5 da aba "P&L FCamara's" = Receita FCamara's (D..I = jan..jun)
PNL_LINHA5 = {'2026-01': 26170714.74, '2026-02': 26024718.04, '2026-03': 26917795.03,
              '2026-04': 26082039.61, '2026-05': 24903384.50, '2026-06': 25201043.02}
GCB = 'CASAS BAHIA'

df = main._get_nova_base()
b = gerar_xlsx_bytes(df, None)
wb = openpyxl.load_workbook(io.BytesIO(b))
ws = wb['Detalhe por Pessoa']
hdr = [ws.cell(4, j).value for j in range(1, ws.max_column + 1)]
dados = []
for r in range(5, ws.max_row):
    linha = {h: ws.cell(r, j).value for j, h in enumerate(hdr, 1)}
    if str(linha.get('Mês') or '').strip():
        dados.append(linha)
wb.close()
t = pd.DataFrame(dados)
MES_INV = {'jan/26': '2026-01', 'fev/26': '2026-02', 'mar/26': '2026-03',
           'abr/26': '2026-04', 'mai/26': '2026-05', 'jun/26': '2026-06'}
# a coluna Mês virou data de verdade (mmm/aa) em 16/09; o mapa de texto fica de reserva
t['per'] = t['Mês'].map(lambda v: v.strftime('%Y-%m') if hasattr(v, 'strftime')
                        else MES_INV.get(str(v).strip()))
t['Valor'] = pd.to_numeric(t['Valor'], errors='coerce').fillna(0)
print(f'arquivo gerado: {len(t)} linhas | R$ {t["Valor"].sum():,.2f}\n')

print('=== 1. arquivo x P&L (linha "Receita FCamara\'s") ===')
print(f'{"mes":9}{"arquivo":>16}{"P&L":>16}{"dif":>14}{"%":>8}   sem Casas Bahia')
semgcb = t[~t['Cliente'].astype(str).str.upper().str.contains(GCB, na=False)]
for m in MESES:
    a = t.loc[t['per'] == m, 'Valor'].sum()
    p = PNL_LINHA5[m]
    s = semgcb.loc[semgcb['per'] == m, 'Valor'].sum()
    print(f'{m[-2:]:9}{a:>16,.0f}{p:>16,.0f}{a - p:>14,.0f}{(a - p) / p * 100:>7.1f}%{s:>18,.0f}')
ta, tp = t['Valor'].sum(), sum(PNL_LINHA5.values())
print(f'{"TOTAL":9}{ta:>16,.0f}{tp:>16,.0f}{ta - tp:>14,.0f}{(ta - tp) / tp * 100:>7.1f}%')

print('\n=== 2. coerencia entre os meses — receita por BU ===')
piv = t.pivot_table(index='BU', columns='per', values='Valor', aggfunc='sum', fill_value=0)
piv = piv.reindex(columns=MESES, fill_value=0)
print(piv.round(0).to_string())
print('\n   variacao m/m (%), sinalizando saltos acima de 25%:')
for bu in piv.index:
    v = piv.loc[bu]
    alertas = []
    for i in range(1, len(MESES)):
        ant, atual = v[MESES[i - 1]], v[MESES[i]]
        if ant and abs(atual - ant) / abs(ant) > 0.25:
            alertas.append(f'{MESES[i][-2:]}: {(atual - ant) / abs(ant) * 100:+.0f}%')
        elif not ant and atual:
            alertas.append(f'{MESES[i][-2:]}: aparece do nada')
        elif ant and not atual:
            alertas.append(f'{MESES[i][-2:]}: some')
    if alertas:
        print(f'      {bu[:20]:20} {" | ".join(alertas)}')

print('\n=== 3. centro de lucro: entra/sai entre os meses ===')
pc = t.pivot_table(index='Centro de Lucro', columns='per', values='Valor', aggfunc='sum', fill_value=0)
pc = pc.reindex(columns=MESES, fill_value=0)
for cod in pc.index:
    v = pc.loc[cod]
    q1 = sum(v[m] for m in MESES[:3])
    q2 = sum(v[m] for m in MESES[3:])
    if (q1 and not q2) or (q2 and not q1):
        print(f'   {cod:10} Q1 R$ {q1:>13,.0f} | Q2 R$ {q2:>13,.0f}   <<< so aparece em um trimestre')

print('\n=== 4. clientes que somem ou aparecem do Q1 para o Q2 (top 8 de cada) ===')
cli = t.pivot_table(index='Cliente', columns='per', values='Valor', aggfunc='sum', fill_value=0)
cli = cli.reindex(columns=MESES, fill_value=0)
cli['q1'] = cli[MESES[:3]].sum(axis=1)
cli['q2'] = cli[MESES[3:]].sum(axis=1)
sumiu = cli[(cli['q1'] > 0) & (cli['q2'] == 0)].sort_values('q1', ascending=False)
surgiu = cli[(cli['q1'] == 0) & (cli['q2'] > 0)].sort_values('q2', ascending=False)
print(f'   somem no Q2: {len(sumiu)} clientes, R$ {sumiu["q1"].sum():,.0f}')
for k, v in sumiu['q1'].head(8).items():
    print(f'      {str(k)[:34]:34} R$ {v:>12,.0f}')
print(f'   so no Q2: {len(surgiu)} clientes, R$ {surgiu["q2"].sum():,.0f}')
for k, v in surgiu['q2'].head(8).items():
    print(f'      {str(k)[:34]:34} R$ {v:>12,.0f}')
