# -*- coding: utf-8 -*-
"""Acrescenta/atualiza os 3 PEPs BR07 da Odontoprev na aba pep_vertical -> Hyper."""
import shutil, sys
import openpyxl
import pandas as pd
P = 'parametros.xlsx'
ALVO = {'BR07CLP00022': 'Hyper', 'BR07CLP00087': 'Hyper', 'BR07CLP00151': 'Hyper'}
d = pd.read_excel(P, sheet_name='pep_vertical')
d.columns = [str(c).strip() for c in d.columns]
atual = {str(r['pep']).strip().upper(): str(r['vertical']).strip() for _, r in d.iterrows()}
print('situacao atual dos 3 PEPs na aba pep_vertical:')
for k in ALVO:
    print(f'   {k}: {atual.get(k, "(nao esta na aba)")}')
if '--apply' not in sys.argv:
    print('\n--- DRY RUN ---'); raise SystemExit(0)
shutil.copy2(P, P.replace('.xlsx', '_antes_odonto_18.09.xlsx'))
wb = openpyxl.load_workbook(P)
ws = wb['pep_vertical']
col = {str(c.value).strip(): c.column for c in ws[1] if c.value}
cp, cv = col['pep'], col['vertical']
existentes = {}
for r in range(2, ws.max_row + 1):
    v = ws.cell(r, cp).value
    if v: existentes[str(v).strip().upper()] = r
n_add = n_upd = 0
for pep, vert in ALVO.items():
    if pep in existentes:
        ws.cell(existentes[pep], cv).value = vert; n_upd += 1
    else:
        r = ws.max_row + 1
        ws.cell(r, cp).value = pep; ws.cell(r, cv).value = vert; n_add += 1
wb.save(P); wb.close()
print(f'\ngravado: {n_add} novos, {n_upd} atualizados (backup: parametros_antes_odonto_18.09.xlsx)')
t = pd.read_excel(P, sheet_name='pep_vertical')
print(f'teste de abertura OK: {len(t)} linhas')
print(t[t['pep'].astype(str).str.upper().isin(ALVO)].to_string(index=False))
