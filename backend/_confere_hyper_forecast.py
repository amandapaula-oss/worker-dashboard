# -*- coding: utf-8 -*-
"""Confere a receita do Hyper na nossa base contra a planilha oficial da carteira
('5. Forecast\\Hyper\\Hyper - Receita x Caixa.xlsx', aba RECEITA LÍQUIDA), cliente a
cliente e mes a mes.

A comparacao soma a receita das fontes 'Hyper' (Q1) e 'Hyper Q2' — a carteira Hyper —
mesmo quando a linha foi reclassificada para outra BU pelo cadastro de clientes.
"""
import os, re, unicodedata
from collections import defaultdict

import openpyxl
import pandas as pd

os.environ.setdefault('SECRET_KEY', 'dbg-readonly')
os.environ.setdefault('ALLOWED_ORIGINS', '*')
from _supabase_creds import load_creds
load_creds()
import main

FC = (r'C:\Users\amanda.paula\FCamara Consultoria e Formação\FCamara Files - CONTROLADORIA'
      r'\30. FP&A NOVO\03. Apresentações\2026\5. Comitê de Finanças - Fechamentos FP&A'
      r'\5. Forecast\Hyper\Hyper - Receita x Caixa.xlsx')
MESES = ['2026-01', '2026-02', '2026-03', '2026-04', '2026-05', '2026-06']


# o forecast usa o apelido comercial; a base usa a razao social
DEPARA = {'BANCO BV': 'VOTORANTIM', 'BV': 'VOTORANTIM',
          'BANCO DE TOKYO MITSUBISHI UFJ BRASIL': 'MUFG BRASIL', 'TOKYO MITSUBISHI': 'MUFG BRASIL',
          'DEXCO': 'DURATEX', 'BANCO OURINVEST': 'OURIBANK', 'OURINVEST': 'OURIBANK',
          'MERCADO LIVRE MERCADO PAGO': 'MERCADO LIVRE', 'CALLINK SERVICOS DE CALL CENTER': 'CALLINK',
          'CIRION': 'CIRION TECHNOLOGIES', 'GRU AIRPORT': 'CONCESSIONARIA DO AEROPORTO',
          'ULTRA': 'ULTRA', 'GRUPO ULTRA': 'ULTRA'}


def npn(v):
    v = unicodedata.normalize('NFKD', str(v or '')).encode('ascii', 'ignore').decode().upper()
    v = re.sub(r'\b(S\.?\s?A\.?|S/?A|LTDA|EIRELI|ME|EPP|SA|GRUPO|BANCO|FL 01|INC|DO BRASIL)\b', ' ', v)
    v = re.sub(r'\(.*?\)', ' ', v)
    v = re.sub(r'\s*-\s*[A-Z/]{2,}\s*$', '', v)
    v = ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', v).split())
    return DEPARA.get(v, v)


# ---------- forecast ----------
wb = openpyxl.load_workbook(FC, data_only=True)
ws = wb['RECEITA LÍQUIDA']
# o cabecalho traz o mes como TEXTO em ingles ("Jan", "Feb"...), nao como data
_EN = {'JAN': '01', 'FEB': '02', 'MAR': '03', 'APR': '04', 'MAY': '05', 'JUN': '06',
       'JUL': '07', 'AUG': '08', 'SEP': '09', 'OCT': '10', 'NOV': '11', 'DEC': '12'}
cols = {}
for linha in (2, 3, 1):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(linha, c).value
        k = None
        if hasattr(v, 'strftime'):
            k = v.strftime('%Y-%m')
        elif isinstance(v, str) and v.strip()[:3].upper() in _EN:
            k = f'2026-{_EN[v.strip()[:3].upper()]}'
        if k in MESES and k not in cols:
            cols[k] = c
    if len(cols) >= 6:
        break
fc = {}
for r in range(4, ws.max_row + 1):
    nome = ws.cell(r, 2).value
    if not nome or str(nome).strip().lower().startswith('total'):
        continue
    vals = {m: float(ws.cell(r, c).value or 0) for m, c in cols.items()}
    if any(vals.values()):
        fc[str(nome).strip()] = vals
wb.close()
print(f'forecast Hyper: {len(fc)} clientes | colunas de mes: {sorted(cols)}')
tot_fc = {m: sum(v.get(m, 0) for v in fc.values()) for m in MESES}
print('  total por mes:', {m: round(x) for m, x in tot_fc.items()})

# ---------- nossa base: a carteira Hyper ----------
df = main._get_nova_base()
d = df[df['periodo'].astype(str).isin(MESES)].copy()
d['receita'] = pd.to_numeric(d['receita'], errors='coerce').fillna(0)
hy = d[d['fonte'].astype(str).isin(['Hyper', 'Hyper Q2']) & (d['receita'] != 0)]
base = defaultdict(lambda: defaultdict(float))
for _, x in hy.iterrows():
    base[npn(x['nome_cliente'])][str(x['periodo'])] += x['receita']
tot_base = {m: sum(v.get(m, 0) for v in base.values()) for m in MESES}
print(f'\nnossa base (fontes Hyper + Hyper Q2): {len(base)} clientes')
print('  total por mes:', {m: round(x) for m, x in tot_base.items()})
print('  diferenca   :', {m: round(tot_base[m] - tot_fc[m]) for m in MESES})

# ---------- cliente a cliente ----------
print('\n=== cliente a cliente (diferenca > R$ 100 em algum mes) ===')
print(f'{"cliente":34} ' + ' '.join(f'{m[-2:]:>11}' for m in MESES))
achou_base = set()
linhas = []
for nome, vals in sorted(fc.items(), key=lambda t: -sum(t[1].values())):
    k = npn(nome)
    b = base.get(k)
    if b is None:
        cand = [c for c in base if c.startswith(k[:11]) or k.startswith(c[:11])]
        if len(cand) == 1:
            k, b = cand[0], base[cand[0]]
    if b is not None:
        achou_base.add(k)
    difs = {m: (b.get(m, 0) if b else 0) - vals.get(m, 0) for m in MESES}
    if any(abs(x) > 100 for x in difs.values()):
        linhas.append((nome, vals, b or {}, difs))
for nome, vals, b, difs in linhas[:24]:
    print(f'{nome[:34]:34} ' + ' '.join(f'{difs[m]:>11,.0f}' for m in MESES))
print(f'   ({len(linhas)} clientes com diferenca)')

sobra = [k for k in base if k not in achou_base]
if sobra:
    print('\n=== na nossa base e NAO no forecast ===')
    for k in sorted(sobra, key=lambda x: -sum(base[x].values()))[:12]:
        print(f'   {k[:34]:34} R$ {sum(base[k].values()):>12,.0f}')
