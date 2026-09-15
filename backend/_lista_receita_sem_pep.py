# -*- coding: utf-8 -*-
"""Gera a planilha das receitas do Q2 que ficaram sem PEP, pra Amanda revisar."""
import os
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

os.environ.setdefault('SECRET_KEY', 'dbg-readonly')
from _supabase_creds import load_creds
import httpx

url, key = load_creds()
H = {'apikey': key, 'Authorization': f'Bearer {key}'}

rows, ids, off = [], set(), 0
while True:
    r = httpx.get(f'{url}/rest/v1/nova_base',
                  params={'select': 'id,fonte,fonte_dados,periodo,empresa,vertical,pep,receita,'
                                    'nome_cliente,nome_pessoa,tipos,no_hierarquia,apuracao',
                          'periodo': 'in.(2026-04,2026-05,2026-06)', 'order': 'id',
                          'limit': '1000', 'offset': str(off)}, headers=H, timeout=90)
    b = r.json()
    if not isinstance(b, list):
        raise SystemExit(str(b)[:200])
    for x in b:
        if x['id'] not in ids:
            ids.add(x['id']); rows.append(x)
    if len(b) < 1000:
        break
    off += 1000

sem = [x for x in rows if (x['receita'] or 0) != 0 and x['fonte'] != 'Budget' and not x['pep']]
TEM_PEP_NA_ORIGEM = {'Time & Expenses'}


def grupo(x):
    t = str(x.get('tipos') or '').strip()
    if t in TEM_PEP_NA_ORIGEM:
        return 'A. DEVERIA ter PEP (alocação T&E)'
    if t in ('Fee', 'WIP', 'Usage Based', 'Ecossistema'):
        return f'B. Sem PEP por natureza ({t})'
    return 'C. Verificar (sem tipo)'


sem.sort(key=lambda x: (grupo(x), -(x['receita'] or 0)))
MES = {'2026-04': 'abr/26', '2026-05': 'mai/26', '2026-06': 'jun/26'}

wb = Workbook()
ws = wb.active
ws.title = 'Receitas sem PEP'
hdr_fill = PatternFill('solid', fgColor='C04F15')
hdr_font = Font(name='Aptos Narrow', size=11, bold=True, color='FFFFFF')
base = Font(name='Aptos Narrow', size=11)
tot_fill = PatternFill('solid', fgColor='FFFFCC')
grp_fill = PatternFill('solid', fgColor='EEEEEE')
thin = Border(bottom=Side(style='thin', color='DDDDDD'))

ws['A1'] = 'Receitas do 2º trimestre/2026 sem elemento PEP — revisar com o time do racional'
ws['A1'].font = Font(name='Aptos Narrow', size=12, bold=True)
ws['A2'] = ('Base: nova_base do dashboard (exclui Budget). Grupo A = alocação por pessoa, que normalmente TEM projeto '
            'no racional — são os casos a investigar. Grupo B = Fee/WIP/Usage Based/Ecossistema, faturamento por '
            'contrato, sem elemento PEP na origem.')
ws['A2'].font = Font(name='Aptos Narrow', size=9, italic=True, color='666666')

cols = ['Mês', 'Empresa', 'BU', 'Cliente', 'Pessoa', 'Tipo', 'Hierarquia', 'Apuração', 'Receita (R$)', 'Fonte']
for j, c in enumerate(cols, 1):
    cell = ws.cell(4, j, c)
    cell.font = hdr_font; cell.fill = hdr_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

r = 5
g_atual = None
for x in sem:
    g = grupo(x)
    if g != g_atual:
        g_atual = g
        tot_g = sum(y['receita'] for y in sem if grupo(y) == g)
        n_g = sum(1 for y in sem if grupo(y) == g)
        c = ws.cell(r, 1, f'{g}  —  {n_g} linhas, R$ {tot_g:,.0f}'.replace(',', '.'))
        c.font = Font(name='Aptos Narrow', size=11, bold=True)
        for j in range(1, len(cols) + 1):
            ws.cell(r, j).fill = grp_fill
        r += 1
    vals = [MES.get(x['periodo'], x['periodo']), x.get('empresa'), x.get('vertical'),
            x.get('nome_cliente'), x.get('nome_pessoa'), x.get('tipos'),
            x.get('no_hierarquia'), x.get('apuracao'), round(float(x['receita']), 2), x.get('fonte')]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(r, j, v)
        cell.font = base
        cell.border = thin
        if j == 9:
            cell.number_format = '#,##0'
    r += 1

ws.cell(r, 1, 'TOTAL').font = Font(name='Aptos Narrow', size=11, bold=True)
c = ws.cell(r, 9, round(sum(x['receita'] for x in sem), 2))
c.font = Font(name='Aptos Narrow', size=11, bold=True)
c.number_format = '#,##0'
for j in range(1, len(cols) + 1):
    ws.cell(r, j).fill = tot_fill

larg = [9, 15, 16, 34, 32, 16, 24, 13, 15, 18]
from openpyxl.utils import get_column_letter
for j, w in enumerate(larg, 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = 'A5'
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 85

out = r'C:\Users\amanda.paula\FCamara Consultoria e Formação\FCamara Files - CONTROLADORIA\30. FP&A NOVO\06. Cockpit - Desenvolvimentos\NewDashboard\Apuração de metas 2 Q\Receitas sem PEP - Q2Y26.xlsx'
wb.save(out)
import shutil
shutil.copy(out, '_receitas_sem_pep_q2.xlsx')
from openpyxl import load_workbook
load_workbook(out)      # testa abertura
print('salvo:', out)
print(f'linhas: {len(sem)} | total R$ {sum(x["receita"] for x in sem):,.0f}')
for g, n in Counter(grupo(x) for x in sem).most_common():
    v = sum(x['receita'] for x in sem if grupo(x) == g)
    print(f'   {g:40} {n:>3} linhas  R$ {v:>12,.0f}')
