# -*- coding: utf-8 -*-
"""Extrai o Q2 das abas de receita que sumiram da base quando a Base Unificada
substituiu a carga de 02/09. Cada aba tem um layout proprio (mapeado antes):

  Licensing Hyper : so existe o TOTAL no Q2 (r14) — nao ha detalhe por cliente
  Play            : detalhe por projeto (col E), meses nas colunas M,N,O
  Sales Boost     : detalhe por cliente (col E), meses nas colunas I,J,K; a r7 e
                    SUBTOTAL do Adcos (= r8+r9+r10) e nao pode entrar
  Licensing Msf   : um bloco de 6 linhas por mes; 3 clientes + TOTAIS; a receita
                    liquida esta na coluna H

Cada aba e conferida contra o proprio total antes de virar carga.
Saida: _abas_faltantes_q2.json
"""
import json
import os
import re
import unicodedata
from collections import defaultdict

import openpyxl
import pandas as pd

os.environ.setdefault('SECRET_KEY', 'dbg-readonly')
os.environ.setdefault('ALLOWED_ORIGINS', '*')
from _supabase_creds import load_creds
load_creds()
import main

DIR = os.path.dirname(os.path.abspath(__file__))
PL = os.path.join(DIR, '2026 dados', 'FCamara - P&L Gerencial - jun26 v2.xlsx')
Q2 = ['2026-04', '2026-05', '2026-06']


def npn(v):
    v = unicodedata.normalize('NFKD', str(v or '')).encode('ascii', 'ignore').decode().upper()
    v = re.sub(r'\b(S\.?\s?A\.?|S/?A|LTDA|EIRELI|ME|EPP|SA|GRUPO|BANCO)\b', ' ', v)
    return ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', v).split())


def num(v):
    n = pd.to_numeric(v, errors='coerce')
    return 0.0 if pd.isna(n) else float(n)


wb = openpyxl.load_workbook(PL, data_only=True)
linhas = []

# ---------- 1. Licensing Hyper: so o total (r14), colunas F,G,H ----------
ws = wb['Licensing Hyper']
for per, c in zip(Q2, (6, 7, 8)):
    v = num(ws.cell(14, c).value)
    if v:
        linhas.append({'fonte': 'Licensing Hyper Q2', 'periodo': per,
                       'cliente': 'Licensing Hyper (sem detalhe por cliente)',
                       'projeto': '', 'valor': round(v, 2)})
print(f"Licensing Hyper: {sum(x['valor'] for x in linhas if x['fonte'] == 'Licensing Hyper Q2'):,.0f}")

# ---------- 2. Play: detalhe por projeto ----------
ws = wb['Play']
tot_play = defaultdict(float)
n0 = len(linhas)
for r in range(5, 40):
    proj = ws.cell(r, 5).value
    if not proj or str(proj).strip().lower().startswith(('total', 'soma')):
        continue
    for per, c in zip(Q2, (13, 14, 15)):
        v = num(ws.cell(r, c).value)
        if v:
            cli = str(proj).split('-')[0].strip()
            linhas.append({'fonte': 'Play Q2', 'periodo': per, 'cliente': cli,
                           'projeto': str(proj).strip(), 'valor': round(v, 2)})
            tot_play[per] += v
print(f'Play: {len(linhas) - n0} linhas | por mes ' +
      str({k: round(v) for k, v in sorted(tot_play.items())}))
for per, c in zip(Q2, (13, 14, 15)):
    print(f'   confere {per}: extraido {tot_play[per]:>12,.0f} | total da aba (r25) '
          f'{num(ws.cell(25, c).value):>12,.0f}')

# ---------- 3. Sales Boost: detalhe, pulando a r7 (subtotal do Adcos) ----------
ws = wb['Sales Boost']
tot_sb = defaultdict(float)
n0 = len(linhas)
for r in range(8, 21):                       # comeca na 8: a 6 e total e a 7 subtotal
    cli = ws.cell(r, 5).value
    if not cli or str(cli).strip().lower().startswith(('total', 'soma')):
        continue
    for per, c in zip(Q2, (9, 10, 11)):
        v = num(ws.cell(r, c).value)
        if v:
            linhas.append({'fonte': 'Sales Boost Q2', 'periodo': per,
                           'cliente': str(cli).strip(), 'projeto': str(ws.cell(r, 3).value or '').strip(),
                           'valor': round(v, 2)})
            tot_sb[per] += v
print(f'\nSales Boost: {len(linhas) - n0} linhas | por mes ' +
      str({k: round(v) for k, v in sorted(tot_sb.items())}))
for per, c in zip(Q2, (9, 10, 11)):
    print(f'   confere {per}: extraido {tot_sb[per]:>12,.0f} | total da aba (r6) '
          f'{num(ws.cell(6, c).value):>12,.0f}')

# ---------- 4. Licensing Msf: bloco de 6 linhas por mes, receita liquida na col H ----------
ws = wb['Licensing Msf']
tot_msf = defaultdict(float)
n0 = len(linhas)
blocos = {'2026-04': (21, 23), '2026-05': (27, 29), '2026-06': (33, 35)}
for per, (ini, fim) in blocos.items():
    for r in range(ini, fim + 1):
        cli = ws.cell(r, 2).value
        if not cli or str(cli).strip().upper().startswith('TOTAIS'):
            continue
        v = num(ws.cell(r, 8).value)
        if v:
            linhas.append({'fonte': 'Licensing Msf Q2', 'periodo': per,
                           'cliente': str(cli).strip(), 'projeto': '', 'valor': round(v, 2)})
            tot_msf[per] += v
print(f'\nLicensing Msf: {len(linhas) - n0} linhas | por mes ' +
      str({k: round(v) for k, v in sorted(tot_msf.items())}))
for per, lt in (('2026-04', 24), ('2026-05', 30), ('2026-06', 36)):
    print(f'   confere {per}: extraido {tot_msf[per]:>12,.0f} | TOTAIS da aba (r{lt}) '
          f'{num(ws.cell(lt, 8).value):>12,.0f}')
wb.close()

# ---------- sobreposicao com a base ----------
df = main._get_nova_base()
d = df[df['periodo'].astype(str).isin(Q2)].copy()
d['receita'] = pd.to_numeric(d['receita'], errors='coerce').fillna(0)
d = d[(d['receita'] != 0) & (d['fonte'].astype(str) != 'Budget')]
ja = defaultdict(float)
for _, x in d.iterrows():
    ja[(npn(x['nome_cliente']), str(x['periodo']))] += x['receita']

print('\n=== risco de duplicar ===')
risco = 0.0
for x in linhas:
    k = (npn(x['cliente']), x['periodo'])
    if k in ja:
        risco += x['valor']
        print(f"   {x['cliente'][:24]:24} {x['periodo']} aba R$ {x['valor']:>10,.0f} | "
              f"base ja tem R$ {ja[k]:>11,.0f}  [{x['fonte']}]")
print(f'   valor sob risco: R$ {risco:,.0f} de R$ {sum(x["valor"] for x in linhas):,.0f}')

json.dump(linhas, open(os.path.join(DIR, '_abas_faltantes_q2.json'), 'w', encoding='utf-8'),
          ensure_ascii=False, indent=1)
print(f'\ntotal a subir: {len(linhas)} linhas | R$ {sum(x["valor"] for x in linhas):,.0f}')
print('salvo _abas_faltantes_q2.json')
