# -*- coding: utf-8 -*-
"""O Licensing Hyper subiu como 1 linha por mes ('sem detalhe por cliente') porque a
area principal da aba esta vazia no Q2. O detalhe existe: fica num bloco lateral da
propria aba (linhas 67-83, colunas R:AD) — e de la que as formulas F14/G14/H14 somam.
Troca as 3 linhas genericas por 11 com cliente, PEP e BU. So grava se os 3 meses
baterem com o total oficial da aba.
Uso: python _refaz_licensing_hyper.py [--apply]
"""
import json, os, sys
import openpyxl

APPLY = '--apply' in sys.argv
DIR = os.path.dirname(os.path.abspath(__file__))
os.environ.setdefault('SECRET_KEY', 'dbg-readonly')
from _supabase_creds import load_creds
url, key = load_creds()
import httpx

H = {'apikey': key, 'Authorization': f'Bearer {key}', 'Content-Type': 'application/json'}
PL = os.path.join(DIR, '2026 dados', 'FCamara - P&L Gerencial - jun26 v2.xlsx')
FONTE = 'Licensing Hyper Q2'
# (periodo, linha inicial, linha final, coluna do valor liquido, coluna do PEP)
BLOCOS = [('2026-04', 67, 71, 26, 22), ('2026-05', 75, 77, 26, 22), ('2026-06', 81, 83, 26, 21)]
OFICIAL = {'2026-04': 1939374.81, '2026-05': 1036385.58, '2026-06': 640362.16}
KEYS = ['fonte', 'fonte_dados', 'periodo', 'empresa', 'pep', 'pep_base', 'nome_pessoa',
        'nome_cliente', 'vertical', 'apuracao_manual', 'tipos', 'receita', 'custo_rateado',
        'classificacao', 'no_hierarquia']

wb = openpyxl.load_workbook(PL, data_only=True)
ws = wb['Licensing Hyper']
rows, tot = [], {}
for per, ini, fim, cval, cpep in BLOCOS:
    for r in range(ini, fim + 1):
        try:
            v = float(ws.cell(r, cval).value)
        except (TypeError, ValueError):
            continue
        if not v:
            continue
        cli = str(ws.cell(r, 20).value or '').strip()
        pep = str(ws.cell(r, cpep).value or '').strip().upper()
        vert = str(ws.cell(r, 29).value or '').strip()
        x = {k: None for k in KEYS}
        x.update({'fonte': FONTE,
                  'fonte_dados': ('FCamara - P&L Gerencial - jun26 v2.xlsx | aba Licensing Hyper '
                                  f'(detalhe lateral, linha {r})'),
                  'periodo': per, 'empresa': 'BR02 FCamara',
                  'pep': pep or None, 'pep_base': pep.split('.')[0] if pep else None,
                  'nome_cliente': cli or 'Licensing Hyper',
                  'vertical': vert or 'BU Hyper', 'apuracao_manual': 'Ecossistema',
                  'no_hierarquia': 'DC009 Licensing Hyper',
                  'tipos': str(ws.cell(r, 21).value or '').strip() or None,
                  'receita': round(v, 2)})
        rows.append(x)
        tot[per] = tot.get(per, 0) + v
wb.close()

print(f'linhas extraidas: {len(rows)}')
for per in OFICIAL:
    print(f'   {per}: extraido R$ {tot.get(per, 0):>12,.2f} | oficial R$ {OFICIAL[per]:>12,.2f} | '
          f'dif {tot.get(per, 0) - OFICIAL[per]:>7,.2f}')
for x in rows:
    print(f"   {x['periodo']} {str(x['nome_cliente'])[:30]:30} {str(x['pep']):20} "
          f"{str(x['vertical']):12} R$ {x['receita']:>12,.2f}")
if any(abs(tot.get(p, 0) - OFICIAL[p]) > 1 for p in OFICIAL):
    raise SystemExit('!! nao bate com o oficial — nao gravo')

if not APPLY:
    print('\n--- DRY RUN: nada gravado. Rode com --apply. ---')
    raise SystemExit(0)

r = httpx.delete(f'{url}/rest/v1/nova_base', params={'fonte': f'eq.{FONTE}'},
                 headers={**H, 'Prefer': 'count=exact'}, timeout=300)
assert r.status_code in (200, 204), r.text[:200]
print(f'\nlinhas genericas removidas: {r.headers.get("content-range")}')
rr = httpx.post(f'{url}/rest/v1/nova_base', json=rows,
                headers={**H, 'Prefer': 'return=representation'}, timeout=300)
assert rr.status_code in (200, 201), f'{rr.status_code} {rr.text[:300]}'
ids = [z['id'] for z in rr.json()]
json.dump(ids, open(os.path.join(DIR, '_backup_insert_licensing_hyper_detalhe.json'), 'w'))
print(f'inseridas: {len(ids)} linhas com cliente e PEP')
